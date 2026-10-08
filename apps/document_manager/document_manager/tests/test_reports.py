# -*- coding: utf-8 -*-
"""Reports of module 4 and the reader reports of module 6, their CSV/XLSX/print output, and the inventory of the fonds.

Run: bench --site <site> run-tests --app document_manager --module document_manager.tests.test_reports
"""

import io
from datetime import date
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import add_days, getdate, nowdate

from document_manager.document_manager.api import inventory
from document_manager.document_manager.api import reports as api
from document_manager.document_manager.services import reports
from document_manager.document_manager.services.reports import archive as archive_reports
from document_manager.document_manager.services.reports import output, readers as reader_reports
from document_manager.document_manager.setup import install
from document_manager.tests.test_request_workflow import _names, _reader, _seed_archive, _user

TAG = frappe.generate_hash(length=5).lower()
ADMIN = f"rp.admin.{TAG}@example.com"
CATALOGER = f"rp.cataloger.{TAG}@example.com"
OFFICER = f"rp.officer.{TAG}@example.com"
LEADER = f"rp.leader.{TAG}@example.com"
PRESERVER = f"rp.preserver.{TAG}@example.com"
READER = f"rp.reader.{TAG}@example.com"
SECOND_READER = f"rp.second.{TAG}@example.com"
USERS = (ADMIN, CATALOGER, OFFICER, LEADER, PRESERVER, READER, SECOND_READER)


def _run(slug, filters=None, **kw):
    return reports.run(slug, filters or {}, kw.get("page", 1), kw.get("page_size", 50))


class TestReports(IntegrationTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        frappe.set_user("Administrator")
        install.seed_all()
        _user(ADMIN, ["Document Admin"], "System User")
        _user(CATALOGER, ["Cataloger"], "System User")
        _user(OFFICER, ["Reading Room Officer"], "System User")
        _user(LEADER, ["Archive Leader"], "System User")
        _user(PRESERVER, ["Preservation Officer"], "System User")
        _user(READER, ["Reader"], "Website User")
        _user(SECOND_READER, ["Reader"], "Website User")
        cls.profile = _reader(READER, "RP Người đọc")
        cls.second = _reader(SECOND_READER, "RP Người đọc hai")
        cls._seeded = _seed_archive(3)
        cls.fonds = _names(cls._seeded, "Fonds")[0]
        cls.catalog = _names(cls._seeded, "Catalog")[0]
        cls.group = _names(cls._seeded, "Record Group")[0]
        cls.files = _names(cls._seeded, "Archival File")
        cls.documents = _names(cls._seeded, "Archive Document")
        cls.agency = _names(cls._seeded, "Archival Agency")[0]
        frappe.db.set_value("Fonds", cls.fonds, {"fonds_code": f"F-{TAG}", "start_year": 1990, "end_year": 2005, "total_boxes": 7})
        frappe.db.set_value("Archival File", cls.files[0], {
            "total_pages": 10, "start_date": "1995-01-10", "shelf_number": "G1", "box_number": "H2", "file_number": "01"})
        frappe.db.set_value("Archival File", cls.files[1], {"total_pages": 5, "start_date": "1996-03-01", "status": "Đang xử lý"})
        frappe.db.set_value("Archival File", cls.files[2], {"total_pages": 0, "start_date": "2001-06-01", "file_number": "03"})
        frappe.db.set_value("Archive Document", cls.documents[0], {
            "page_count": 4, "file_size_kb": 2048, "document_date": "1995-02-02", "author": "Bộ Nội vụ", "document_number": "12/QĐ",
            "file_type": "PDF"})
        frappe.db.set_value("Archive Document", cls.documents[1], {"page_count": 6, "file_size_kb": 1024, "document_date": "1996-04-04",
                                                                   "author": "UBND tỉnh", "file_type": "DOCX"})
        frappe.db.commit()

    @classmethod
    def tearDownClass(cls):
        frappe.set_user("Administrator")
        for dt in ("Usage Request", "Copy Request"):
            for name in frappe.get_all(dt, filters={"reader": ["in", [cls.profile, cls.second]]}, pluck="name"):
                frappe.delete_doc(dt, name, force=True, ignore_permissions=True)
        for dt, name in cls._seeded:
            frappe.delete_doc(dt, name, force=True, ignore_permissions=True)
        for profile in (cls.profile, cls.second):
            frappe.delete_doc("Reader", profile, force=True, ignore_permissions=True)
        for email in USERS:
            frappe.delete_doc("User", email, force=True, ignore_permissions=True)
        frappe.db.commit()
        super().tearDownClass()

    def setUp(self):
        commit = patch.object(frappe.db, "commit")
        commit.start()
        self.addCleanup(commit.stop)
        frappe.set_user(ADMIN)
        frappe.db.savepoint("rp_test")

    def tearDown(self):
        frappe.set_user("Administrator")
        frappe.db.rollback(save_point="rp_test")

    # ---- the registry and who may run what
    def test_every_report_runs_with_no_filter_and_returns_what_it_declares(self):
        for slug, report in reports.REPORTS.items():
            result = _run(slug)
            self.assertEqual(result["slug"], slug)
            self.assertIsInstance(result["rows"], list, slug)
            self.assertGreaterEqual(result["total"], len(result["rows"]), slug)
            wanted = {c["fieldname"] for c in result["columns"]}
            for row in result["rows"][:5]:
                self.assertTrue(wanted <= set(row) | {"name"}, f"{slug}: missing {wanted - set(row)}")
            self.assertTrue(report.title and report.group and report.icon, slug)

    def test_the_catalogue_of_reports_follows_the_users_rights(self):
        def slugs(user):
            frappe.set_user(user)
            return {r["slug"] for g in api.list_reports() for r in g["reports"]}

        everything = set(reports.REPORTS)
        self.assertEqual(slugs(ADMIN), everything)
        self.assertEqual(slugs(LEADER), everything)
        self.assertEqual(slugs(OFFICER), everything - {"tong-kiem-ke"})
        archive_only = slugs(CATALOGER)
        self.assertIn("phong", archive_only)
        self.assertNotIn("doc-gia", archive_only)
        self.assertNotIn("phieu-qua-han", archive_only)
        self.assertEqual(slugs(PRESERVER), set())

    def test_readers_and_guests_cannot_run_a_report(self):
        for user in (READER, SECOND_READER, "Guest"):
            frappe.set_user(user)
            with self.assertRaises(frappe.PermissionError, msg=user):
                api.list_reports()
            with self.assertRaises(frappe.PermissionError, msg=user):
                api.run_report("phong")
            with self.assertRaises(frappe.PermissionError, msg=user):
                api.get_report("doc-gia")
            with self.assertRaises(frappe.PermissionError, msg=user):
                api.download_report("phong")
        frappe.set_user(CATALOGER)
        with self.assertRaises(frappe.PermissionError):
            api.run_report("doc-gia")
        with self.assertRaises(frappe.DoesNotExistError):
            api.run_report("khong-co")

    def test_filters_are_typed_and_validated(self):
        with self.assertRaises(frappe.ValidationError):
            _run("ho-so", {"status": "Không có trạng thái này"})
        with self.assertRaises(frappe.ValidationError):
            _run("ho-so", {"date_from": "không phải ngày"})
        ignored = _run("ho-so", {"fonds": self.fonds, "nonsense": "x"})
        self.assertNotIn("nonsense", ignored["filters"])
        self.assertEqual(ignored["total"], 3)
        self.assertEqual(_run("phong", {"year_from": "1991"})["filters"]["year_from"], 1991)

    # ---- the catalogue reports
    def test_fonds_report_counts_every_level(self):
        result = _run("phong", {"q": f"F-{TAG}"})
        self.assertEqual(result["total"], 1)
        row = result["rows"][0]
        self.assertEqual((row["record_groups"], row["catalogs"], row["files"], row["documents"]), (1, 1, 3, 3))
        self.assertEqual((row["pages"], row["size_mb"], row["total_boxes"]), (15, 3.0, 7))
        self.assertEqual(row["period"], "1990 – 2005")
        self.assertEqual(row["archival_agency"], self.agency)
        self.assertEqual(result["summary"]["files"], 3)

    def test_fonds_report_filters_by_years_and_agency(self):
        base = {"archival_agency": self.agency}
        self.assertEqual(_run("phong", {**base, "year_from": 2010})["total"], 0)  # the fonds ended in 2005
        self.assertEqual(_run("phong", {**base, "year_to": 1980})["total"], 0)  # and began in 1990
        self.assertEqual(_run("phong", {**base, "year_from": 2000, "year_to": 2001})["total"], 1)
        self.assertEqual(_run("phong", {**base, "status": "Đã đóng"})["total"], 0)

    def test_catalogue_report(self):
        result = _run("muc-luc", {"fonds": self.fonds})
        self.assertEqual(result["total"], 1)
        self.assertEqual((result["rows"][0]["files"], result["rows"][0]["documents"], result["rows"][0]["pages"]), (3, 3, 15))
        self.assertEqual(_run("muc-luc", {"record_group": self.group})["total"], 1)
        self.assertEqual(_run("muc-luc", {"q": "không có mục lục nào như thế"})["total"], 0)

    def test_file_report_filters_pages_and_totals(self):
        everything = _run("ho-so", {"fonds": self.fonds})
        self.assertEqual(everything["total"], 3)
        self.assertEqual(everything["summary"]["total_pages"], 15)
        located = next(r for r in everything["rows"] if r["file_number"] == "01")
        self.assertIn("Giá G1", located["location"])
        self.assertIn("Hộp H2", located["location"])
        self.assertEqual(_run("ho-so", {"fonds": self.fonds, "status": "Đang xử lý"})["total"], 1)
        self.assertEqual(_run("ho-so", {"fonds": self.fonds, "date_from": "1996-01-01", "date_to": "1999-12-31"})["total"], 1)
        self.assertEqual(_run("ho-so", {"fonds": self.fonds, "q": "File 2"})["total"], 1)
        page = _run("ho-so", {"fonds": self.fonds}, page=2, page_size=2)
        self.assertEqual((page["total"], len(page["rows"]), page["page"]), (3, 1, 2))

    def test_document_report_filters_by_author_date_and_type(self):
        base = {"fonds": self.fonds}
        self.assertEqual(_run("van-ban", base)["total"], 3)
        self.assertEqual(_run("van-ban", {**base, "author": "Nội vụ"})["total"], 1)
        self.assertEqual(_run("van-ban", {**base, "date_from": "1996-01-01"})["total"], 1)
        self.assertEqual(_run("van-ban", {**base, "file_type": "DOCX"})["total"], 1)
        self.assertEqual(_run("van-ban", {**base, "q": "12/QĐ"})["total"], 1)
        result = _run("van-ban", {**base, "archival_file": self.files[0]})
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["rows"][0]["size_mb"], 2.0)
        self.assertEqual(result["summary"]["page_count"], 4)

    def test_a_search_text_is_matched_literally(self):
        self.assertEqual(_run("van-ban", {"fonds": self.fonds, "q": "%"})["total"], 0)
        self.assertEqual(_run("ho-so", {"fonds": self.fonds, "q": "_"})["total"], 0)

    # ---- statistics
    def test_statistics_by_fonds_and_the_other_dimensions_agree(self):
        by_fonds = _run("thong-ke-phong", {"fonds": self.fonds})
        self.assertEqual(by_fonds["total"], 1)
        row = by_fonds["rows"][0]
        self.assertEqual((row["files"], row["documents"], row["pages"], row["files_pct"]), (3, 3, 15, 100.0))
        self.assertEqual(by_fonds["chart"]["datasets"][0]["values"], [3])
        for dimension in (archive_reports.BY_YEAR, archive_reports.BY_LEVEL, archive_reports.BY_STATUS):
            other = _run("thong-ke-phong", {"fonds": self.fonds, "group_by": dimension})
            self.assertEqual(other["summary"]["files"], 3, dimension)
            self.assertEqual(sum(r["files"] for r in other["rows"]), 3, dimension)
        by_year = _run("thong-ke-phong", {"fonds": self.fonds, "group_by": archive_reports.BY_YEAR})
        self.assertEqual([r["label"] for r in by_year["rows"]], ["1995", "1996", "2001"])
        by_status = _run("thong-ke-phong", {"fonds": self.fonds, "group_by": archive_reports.BY_STATUS})
        self.assertEqual({r["label"]: r["files"] for r in by_status["rows"]}, {"Đã hoàn thành": 2, "Đang xử lý": 1})

    def test_statistics_by_type_and_warehouse_roll_up_to_the_warehouse(self):
        category = frappe.get_doc({"doctype": "Document Type Category", "type_name": f"RP loại {TAG}"}).insert()
        root = frappe.get_doc({"doctype": "Storage Warehouse", "warehouse_name": f"RP kho {TAG}", "warehouse_type": "Kho", "is_group": 1}).insert()
        shelf = frappe.get_doc({"doctype": "Storage Warehouse", "warehouse_name": f"RP giá {TAG}", "warehouse_type": "Giá",
                                "parent_warehouse": root.name}).insert()
        frappe.db.set_value("Archival File", self.files[0], {"document_type_category": category.name, "storage_warehouse": shelf.name})
        frappe.db.set_value("Archival File", self.files[1], {"document_type_category": category.name, "storage_warehouse": root.name})
        base = {"fonds": self.fonds}
        by_type = _run("thong-ke-loai-hinh-kho", {**base, "group_by": archive_reports.BY_TYPE})
        counts = {r["label"]: r["files"] for r in by_type["rows"]}
        self.assertEqual(counts[category.name], 2)
        self.assertEqual(counts[reports.base.UNKNOWN], 1)
        by_warehouse = _run("thong-ke-loai-hinh-kho", {**base, "group_by": archive_reports.BY_WAREHOUSE})
        self.assertEqual({r["label"]: r["files"] for r in by_warehouse["rows"]}, {root.name: 2, reports.base.UNKNOWN: 1})
        by_place = _run("thong-ke-loai-hinh-kho", {**base, "group_by": archive_reports.BY_PLACE})
        self.assertEqual({r["label"]: r["files"] for r in by_place["rows"]}[shelf.name], 1)

    # ---- reader reports
    def _slip(self, reader, state, doctype="Usage Request", request_date=None, **extra):
        doc = frappe.get_doc({"doctype": doctype, "reader": reader, "request_date": request_date or nowdate(),
                              "purpose": f"RP {TAG}", "items": [{"archival_file": self.files[0]}, {"archival_file": self.files[1]}]})
        doc.flags.ignore_permissions = True
        doc.flags.ignore_validate = True
        doc.insert()
        frappe.db.set_value(doctype, doc.name, {"workflow_state": state, "docstatus": 0 if state == "Nháp" else 1, **extra})
        return doc.name

    def test_reader_report_counts_slips_per_reader(self):
        self._slip(self.profile, "Chờ duyệt")
        self._slip(self.profile, "Đang sử dụng", is_overdue=1, due_date=add_days(nowdate(), -3))
        self._slip(self.profile, "Nháp")  # a basket is not a slip yet
        self._slip(self.profile, "Chờ duyệt", doctype="Copy Request")
        result = _run("doc-gia", {"q": "RP Người đọc"})
        by_name = {r["full_name"]: r for r in result["rows"]}
        mine, other = by_name["RP Người đọc"], by_name["RP Người đọc hai"]
        self.assertEqual((mine["usage_slips"], mine["copy_slips"], mine["overdue"]), (2, 1, 1))
        self.assertEqual((other["usage_slips"], other["copy_slips"], other["overdue"]), (0, 0, 0))
        self.assertEqual(mine["last_request"], getdate(nowdate()))
        self.assertEqual(_run("doc-gia", {"q": "RP Người đọc", "state": "Ngừng hoạt động"})["total"], 0)

    def test_slip_statistics_group_by_state_and_month(self):
        self._slip(self.profile, "Chờ duyệt")
        self._slip(self.profile, "Từ chối")
        self._slip(self.second, "Đã trả", request_date="2020-02-15")
        by_state = _run("thong-ke-phieu", {"group_by": "Trạng thái", "date_from": str(add_days(nowdate(), -1))})
        counts = {r["label"]: r["slips"] for r in by_state["rows"]}
        self.assertEqual(counts["Chờ duyệt"], 1)
        self.assertEqual(counts["Từ chối"], 1)
        self.assertNotIn("Đã trả", counts)  # sent in 2020: outside the dates
        by_month = _run("thong-ke-phieu", {"group_by": "Tháng", "date_to": "2020-12-31", "date_from": "2020-01-01"})
        self.assertEqual({r["label"]: r["slips"] for r in by_month["rows"]}.get("2020-02"), 1)
        self.assertEqual(by_month["summary"]["granted"], 1)
        copy_only = _run("thong-ke-phieu", {"kind": reader_reports.COPY_ONLY})
        self.assertEqual(copy_only["summary"]["slips"], sum(r["slips"] for r in copy_only["rows"]))

    def test_overdue_report_lists_what_is_still_out(self):
        late = self._slip(self.profile, "Đang sử dụng", is_overdue=1, due_date=add_days(nowdate(), -4), renewal_count=1)
        fine = self._slip(self.second, "Đang sử dụng", is_overdue=0, due_date=add_days(nowdate(), 5))
        only_late = _run("phieu-qua-han", {})
        names = {r["name"]: r for r in only_late["rows"]}
        self.assertIn(late, names)
        self.assertNotIn(fine, names)
        self.assertEqual((names[late]["days_late"], names[late]["renewal_count"], names[late]["items"]), (4, 1, 2))
        both = {r["name"] for r in _run("phieu-qua-han", {"include_not_due": 1})["rows"]}
        self.assertTrue({late, fine} <= both)

    def test_most_requested_items(self):
        for _i in range(2):
            self._slip(self.profile, "Đã duyệt")
        self._slip(self.second, "Đã duyệt")
        self._slip(self.second, "Chờ duyệt")  # not granted: not counted
        result = _run("tai-lieu-duoc-khai-thac", {"date_from": nowdate()})
        first = next(r for r in result["rows"] if r["name"] == self.files[0])
        self.assertEqual((first["times"], first["readers"], first["kind"]), (3, 2, "Hồ sơ"))
        self.assertEqual(first["route"], f"/ho-so/{self.files[0]}")
        documents_only = _run("tai-lieu-duoc-khai-thac", {"date_from": nowdate(), "level": reader_reports.LEVEL_DOCS})
        self.assertNotIn(self.files[0], {r["name"] for r in documents_only["rows"]})

    # ---- the inventory of the fonds
    def _check(self, **extra):
        return frappe.get_doc({"doctype": "Inventory Check", "check_title": f"RP kiểm kê {TAG}", "check_date": nowdate(),
                               "fonds": self.fonds, **extra}).insert()

    def test_inventory_takes_the_figures_of_the_system_and_the_count_gives_the_difference(self):
        check = self._check()
        self.assertEqual(check.status, "Nháp")
        inventory.populate(check.name)
        check.reload()
        self.assertEqual(check.status, "Đang kiểm kê")
        self.assertEqual(len(check.items), 1)
        line = check.items[0]
        self.assertEqual((line.recorded_files, line.recorded_documents, line.recorded_boxes), (3, 3, 7))
        self.assertEqual((line.counted_files, line.counted_boxes), (3, 7))  # prefilled: only the differences are typed
        line.counted_files, line.counted_boxes, line.counted = 2, 9, 1
        check.save()
        line = check.items[0]
        self.assertEqual((line.difference_files, line.difference_documents, line.difference_boxes), (-1, 0, 2))
        self.assertEqual(check.total_difference, 1)
        report = _run("tong-kiem-ke", {"inventory_check": check.name})
        self.assertEqual(report["rows"][0]["difference_files"], -1)
        self.assertEqual(report["summary"]["difference_boxes"], 2)
        self.assertEqual(dict(report["info"])["Phạm vi"], frappe.db.get_value("Fonds", self.fonds, "fonds_name"))

    def test_the_recorded_figures_are_not_typed_and_a_new_fonds_line_reads_the_system(self):
        check = self._check()
        inventory.populate(check.name)
        check.reload()
        check.items[0].recorded_files = 99  # whatever a client sends, the figure stays the system's
        check.append("items", {"fonds": self._other_fonds(), "counted": 0})
        check.save()
        self.assertEqual(check.items[0].recorded_files, 3)
        self.assertEqual(check.items[1].recorded_files, 0)
        self.assertEqual(check.items[1].recorded_boxes, 0)

    def _other_fonds(self):
        return frappe.get_doc({"doctype": "Fonds", "fonds_name": f"RP phông khác {TAG}", "archival_agency": self.agency}).insert().name

    def test_a_check_is_completed_only_when_every_fonds_is_counted_then_it_is_locked(self):
        check = self._check()
        inventory.populate(check.name)
        with self.assertRaises(frappe.ValidationError):
            inventory.complete(check.name)  # nothing marked counted yet
        check.reload()
        check.items[0].counted = 1
        check.save()
        done = inventory.complete(check.name)
        self.assertEqual(done["status"], "Hoàn thành")
        check.reload()
        self.assertTrue(check.completed_on)
        check.items[0].counted_files = 50
        with self.assertRaises(frappe.ValidationError):
            check.save()  # the figures are locked
        check.reload()
        check.notes = "Đã đối chiếu với sổ"  # a remark may still be added
        check.save()
        with self.assertRaises(frappe.ValidationError):
            inventory.populate(check.name)

    def test_only_an_administrator_reopens_a_completed_check(self):
        check = self._check()
        inventory.populate(check.name)
        check.reload()
        check.items[0].counted = 1
        check.save()
        inventory.complete(check.name)
        frappe.set_user(CATALOGER)
        with self.assertRaises(frappe.PermissionError):
            inventory.reopen(check.name)
        frappe.set_user(ADMIN)
        self.assertEqual(inventory.reopen(check.name)["status"], "Đang kiểm kê")

    def test_inventory_actions_belong_to_the_cataloguing_roles(self):
        check = self._check()
        for user in (OFFICER, READER, "Guest"):
            frappe.set_user(user)
            with self.assertRaises(frappe.PermissionError, msg=user):
                inventory.populate(check.name)
        frappe.set_user(LEADER)
        with self.assertRaises(frappe.PermissionError):
            inventory.complete(check.name)  # a leader reads the result
        frappe.set_user(CATALOGER)
        self.assertEqual(inventory.populate(check.name)["status"], "Đang kiểm kê")

    def test_duplicate_fonds_lines_and_negative_counts_are_refused(self):
        check = self._check()
        inventory.populate(check.name)
        check.reload()
        check.append("items", {"fonds": self.fonds})
        with self.assertRaises(frappe.ValidationError):
            check.save()
        check.reload()
        check.items[0].counted_documents = -1
        with self.assertRaises(frappe.ValidationError):
            check.save()

    def test_without_a_check_the_inventory_report_says_so(self):
        frappe.db.delete("Inventory Check")
        result = _run("tong-kiem-ke")
        self.assertEqual(result["rows"], [])
        self.assertTrue(result["info"])

    # ---- output
    def test_csv_has_a_bom_the_columns_and_vietnamese_dates(self):
        frappe.set_user(ADMIN)
        api.download_report("ho-so", {"fonds": self.fonds, "date_from": "1995-01-01", "date_to": "1995-12-31"}, "csv")
        response = frappe.local.response
        self.assertEqual(response["type"], "download")
        self.assertTrue(response["filename"].startswith("bao-cao-ho-so-") and response["filename"].endswith(".csv"))
        text = response["filecontent"].decode("utf-8")
        self.assertTrue(text.startswith("﻿"))
        lines = text.lstrip("﻿").splitlines()
        self.assertEqual(lines[0].split(",")[0], "Phông")
        self.assertIn("10/01/1995", lines[1])
        self.assertEqual(len(lines), 3)  # header, the one file in 1995, totals

    def test_xlsx_opens_with_numbers_and_the_unit_header(self):
        from openpyxl import load_workbook

        api.download_report("phong", {"q": f"F-{TAG}"}, "xlsx")
        content = frappe.local.response["filecontent"]
        self.assertEqual(content[:2], b"PK")
        sheet = load_workbook(io.BytesIO(content)).active
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        header = next(i for i, row in enumerate(values) if row[:2] == ["Mã phông", "Tên phông"])
        data = values[header + 1]
        self.assertEqual(data[0], f"F-{TAG}")
        self.assertEqual(data[7], 3)  # files, a number and not a text
        self.assertTrue(any(row and row[0] == "Báo cáo phông lưu trữ" for row in values))

    def test_a_download_larger_than_the_limit_is_refused(self):
        with patch.object(output, "MAX_EXPORT_ROWS", 2):
            with self.assertRaises(frappe.ValidationError):
                api.download_report("ho-so", {"fonds": self.fonds}, "csv")

    def test_the_printout_carries_the_unit_header_the_filters_and_escapes_what_it_prints(self):
        frappe.db.set_single_value("Organization Info", "org_name", "Trung tâm <b>Lưu trữ</b>")
        frappe.db.set_value("Archival File", self.files[2], "file_title", "<script>alert(1)</script> hồ sơ")
        api.print_report("ho-so", {"fonds": self.fonds, "status": "Đã hoàn thành"})
        response = frappe.local.response
        self.assertEqual((response["type"], response["display_content_as"]), ("download", "inline"))
        html = response["filecontent"].decode("utf-8")
        self.assertIn("Trung tâm &lt;b&gt;Lưu trữ&lt;/b&gt;", html)
        self.assertNotIn("<script>alert(1)</script>", html)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt; hồ sơ", html)
        self.assertIn("Tình trạng: <strong>Đã hoàn thành</strong>", html)
        self.assertIn("Người lập biểu", html)

    def test_the_printout_as_a_pdf(self):
        api.print_report("phong", {"q": f"F-{TAG}"}, "pdf")
        response = frappe.local.response
        self.assertEqual(response["type"], "pdf")
        self.assertTrue(response["filecontent"].startswith(b"%PDF"))

    def test_output_is_staff_only_and_unknown_formats_are_refused(self):
        with self.assertRaises(frappe.ValidationError):
            api.download_report("phong", None, "docx")
        with self.assertRaises(frappe.ValidationError):
            api.print_report("phong", None, "exe")
        frappe.set_user(READER)
        with self.assertRaises(frappe.PermissionError):
            api.print_report("phong")

    def test_exports_are_logged(self):
        before = frappe.db.count("Business Activity Log", {"reference_doctype": "Report", "reference_name": "phong"})
        api.download_report("phong", None, "csv")
        self.assertEqual(frappe.db.count("Business Activity Log", {"reference_doctype": "Report", "reference_name": "phong"}), before + 1)

    def test_number_and_date_formats_are_vietnamese(self):
        self.assertEqual(output.vn_number(1234567), "1.234.567")
        self.assertEqual(output.vn_number(1234.5, 2), "1.234,50")
        column = {"fieldtype": "Date"}
        self.assertEqual(output.display(column, date(2026, 10, 7)), "07/10/2026")
        self.assertEqual(output.display({"fieldtype": "Percent"}, 12.5), "12,5%")
        self.assertEqual(output.display({"fieldtype": "Data"}, None), "")
